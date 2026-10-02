#!/usr/bin/env python3
import io
import json
import logging
import os
import re
import uuid
from datetime import datetime, timezone
from email import policy
from email.parser import BytesParser
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

from legal_ai import analyze_document, answer_question_from_document, build_document_template, extract_pdf_text


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "tmp_uploads"
GENERATED_DIR = BASE_DIR / "generated"
MAX_UPLOAD_SIZE = 15 * 1024 * 1024
SESSION = {
    "active_document": None,
    "chat_history": [],
    "generated_documents": {},
}


def validate_pdf_upload(filename: str, content_type: str | None, data: bytes) -> str | None:
  if len(data) > MAX_UPLOAD_SIZE:
    return "too_large"
  if not filename.lower().endswith(".pdf"):
    return "invalid"
  if content_type and content_type.lower().split(";", 1)[0].strip() != "application/pdf":
    return "invalid"
  if not data.startswith(b"%PDF-"):
    return "invalid"
  return None


def ensure_storage_dirs():
    UPLOAD_DIR.mkdir(exist_ok=True)
    GENERATED_DIR.mkdir(exist_ok=True)


def clear_active_document():
    document = SESSION.get("active_document") or {}
    file_path = document.get("file_path")
    if file_path and os.path.exists(file_path):
        try:
            os.remove(file_path)
        except OSError:
            pass
    SESSION["active_document"] = None
    SESSION["chat_history"] = []


def register_generated_document(doc_id: str, payload: dict):
    SESSION["generated_documents"][doc_id] = payload


def _coerce_document_blocks(title: str, content: list[str] | list[dict]) -> list[dict[str, str]]:
    if content and isinstance(content[0], dict):
        return content
    blocks = [{"type": "title", "text": title.upper()}]
    for line in content:
        text = str(line).strip()
        if not text:
            continue
        if re.match(r"^\d+(?:\.\d+)*\.\s+[A-Z][A-Z0-9 &/,'-]+$", text) or text in {"SIGNATURES", "NOTICE", "SUBJECT"}:
            kind = "section_heading"
        elif text.isupper() and len(text) < 70:
            kind = "party_heading"
        elif "SIGNATURE:" in text or "SIGNATURE" == text:
            kind = "signature"
        else:
            kind = "paragraph"
        blocks.append({"type": kind, "text": text})
    return blocks


def build_pdf_document(title: str, body_lines: list[str] | list[dict]) -> bytes:
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="LegalTitle", parent=styles["Title"], fontName="Times-Bold", fontSize=16, leading=20, alignment=1, spaceAfter=20, textColor=colors.HexColor("#192520")))
    styles.add(ParagraphStyle(name="LegalSection", parent=styles["Heading2"], fontName="Times-Bold", fontSize=11, leading=15, spaceBefore=12, spaceAfter=5, keepWithNext=True, textColor=colors.HexColor("#192520")))
    styles.add(ParagraphStyle(name="LegalParty", parent=styles["Heading3"], fontName="Times-Bold", fontSize=10, leading=14, spaceBefore=8, spaceAfter=3, keepWithNext=True, textColor=colors.HexColor("#192520")))
    styles.add(ParagraphStyle(name="LegalBody", parent=styles["BodyText"], fontName="Times-Roman", fontSize=11, leading=16, alignment=4, spaceAfter=8, textColor=colors.HexColor("#252d29")))
    styles.add(ParagraphStyle(name="LegalSignature", parent=styles["BodyText"], fontName="Times-Roman", fontSize=10, leading=17, alignment=0, spaceBefore=4, spaceAfter=3))
    blocks = _coerce_document_blocks(title, body_lines)
    story = []
    for block in blocks:
        text = escape(str(block.get("text", ""))).replace("\n", "<br/>")
        block_type = block.get("type")
        style = {"title": styles["LegalTitle"], "section_heading": styles["LegalSection"], "party_heading": styles["LegalParty"], "signature": styles["LegalSignature"]}.get(block_type, styles["LegalBody"])
        paragraph = Paragraph(text, style)
        if block_type in {"section_heading", "party_heading"}:
            story.append(KeepTogether([paragraph]))
        else:
            story.append(paragraph)
        if block_type == "party_heading":
            story.append(Spacer(1, 2 * mm))

    def add_page_chrome(pdf, document):
        pdf.saveState()
        page_width, _ = A4
        pdf.setStrokeColor(colors.HexColor("#dce3de"))
        pdf.setLineWidth(0.5)
        pdf.line(document.leftMargin, 17 * mm, page_width - document.rightMargin, 17 * mm)
        pdf.setFont("Times-Roman", 8)
        pdf.setFillColor(colors.HexColor("#63716a"))
        pdf.drawString(document.leftMargin, 11 * mm, title.upper())
        pdf.drawRightString(page_width - document.rightMargin, 11 * mm, f"Page {document.page}")
        pdf.restoreState()

    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=25 * mm,
        leftMargin=25 * mm,
        topMargin=24 * mm,
        bottomMargin=25 * mm,
        title=title.upper(),
    )
    document.build(story, onFirstPage=add_page_chrome, onLaterPages=add_page_chrome)
    return buffer.getvalue()


def build_docx_document(title: str, body_lines: list[str] | list[dict]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run(title.upper()).font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    for block in _coerce_document_blocks(title, body_lines):
        text = str(block.get("text", ""))
        kind = block.get("type")
        if kind == "title":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(18)
            run = paragraph.add_run(text.upper())
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
        elif kind == "section_heading":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(11)
        elif kind == "party_heading":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(8)
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
        else:
            paragraph = doc.add_paragraph()
            if kind == "signature":
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for index, line in enumerate(text.split("\n")):
                if index:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


INDEX_HTML = """
<!DOCTYPE html>
<html lang="en">
  <head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>Legal AI Copilot</title>
    <style>
      :root {
        --bg: #f4f7fb;
        --panel: #ffffff;
        --soft: #eef3f9;
        --line: #dfe7f1;
        --nav: #10233d;
        --nav-soft: #1b365a;
        --text: #122033;
        --muted: #5c6d82;
        --primary: #153a67;
        --primary-soft: #eef5ff;
        --accent: #2b70d5;
        --success: #1f8a5b;
        --success-soft: #ebf9f1;
        --warning: #a66a00;
        --warning-soft: #fff6df;
        --danger: #bc3f3f;
        --danger-soft: #fdecec;
        --shadow: 0 10px 24px rgba(17, 24, 39, 0.08);
      }
      * { box-sizing: border-box; }
      [hidden] { display: none !important; }
      html, body { margin: 0; min-height: 100%; font-family: Inter, "Segoe UI", Arial, sans-serif; background: var(--bg); color: var(--text); }
      body { min-height: 100vh; }
      button, input, select, textarea { font: inherit; }
      .app-shell { display: flex; min-height: 100vh; }
      .sidebar {
        width: 260px; background: var(--nav); color: white; display: flex; flex-direction: column; padding: 24px 16px; border-right: 1px solid rgba(255,255,255,0.05);
      }
      .brand {
        font-size: 1.2rem; letter-spacing: 0.08em; font-weight: 800; padding: 10px 12px 22px; border-bottom: 1px solid rgba(255,255,255,0.08); margin-bottom: 20px;
      }
      .nav { display: flex; flex-direction: column; gap: 8px; }
      .nav-btn, .sidebar-bottom button {
        width: 100%; border: none; background: transparent; color: rgba(255,255,255,0.8); text-align: left; font-weight: 600; padding: 12px 14px; border-radius: 10px; cursor: pointer; transition: 0.2s ease; display: flex; align-items: center; gap: 10px;
      }
      .nav-btn:hover, .sidebar-bottom button:hover { background: rgba(255,255,255,0.08); color: white; }
      .nav-btn.active { background: rgba(255,255,255,0.12); color: white; box-shadow: inset 0 0 0 1px rgba(255,255,255,0.08); }
      .sidebar-bottom { margin-top: auto; padding-top: 18px; border-top: 1px solid rgba(255,255,255,0.08); display: flex; flex-direction: column; gap: 8px; }
      .workspace { flex: 1; padding: 28px 28px 40px; }
      .topbar {
        background: rgba(255,255,255,0.7); backdrop-filter: blur(10px); border: 1px solid var(--line); border-radius: 18px; padding: 20px 24px; display: flex; align-items: center; justify-content: space-between; gap: 24px; margin-bottom: 22px;
      }
      .eyebrow {
        margin: 0 0 6px; color: var(--muted); font-size: 0.76rem; letter-spacing: 0.08em; text-transform: uppercase; font-weight: 700;
      }
      .topbar h1 { margin: 0; font-size: clamp(1.8rem, 2vw, 2.4rem); }
      .status-pill {
        background: var(--success-soft); color: var(--success); border: 1px solid rgba(31,138,91,0.18); border-radius: 999px; padding: 8px 12px; font-size: 0.8rem; font-weight: 700;
      }
      .page { display: none; }
      .page.active { display: block; }
      .grid { display: grid; gap: 18px; }
      .cards-2 { grid-template-columns: repeat(2, minmax(0, 1fr)); }
      .cards-3 { grid-template-columns: repeat(3, minmax(0, 1fr)); }
      .card {
        background: var(--panel); border: 1px solid var(--line); border-radius: 18px; box-shadow: var(--shadow); padding: 20px;
      }
      .upload-zone {
        border: 1.5px dashed rgba(21,58,103,0.35); background: linear-gradient(180deg, #f9fbff 0%, #f2f7ff 100%); border-radius: 20px; padding: 42px 16px; text-align: center; min-height: 220px; display: grid; place-items: center; gap: 12px;
      }
      .upload-zone h3 { margin: 0; font-size: 1.7rem; }
      .muted { color: var(--muted); }
      .primary-btn, .secondary-btn, .danger-btn, .small-btn {
        border: none; border-radius: 10px; padding: 11px 18px; cursor: pointer; font-weight: 700; transition: 0.2s ease; text-decoration: none; display: inline-flex; align-items: center; justify-content: center; gap: 8px;
      }
      .primary-btn { background: var(--primary); color: white; }
      .primary-btn:hover { background: var(--nav-soft); }
      .secondary-btn { background: var(--soft); color: var(--text); border: 1px solid var(--line); }
      .danger-btn { background: var(--danger-soft); color: var(--danger); border: 1px solid rgba(188,63,63,0.15); }
      .small-btn { padding: 8px 12px; font-size: 0.82rem; }
      .feature-grid { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 16px; margin-top: 22px; }
      .feature-card { background: var(--panel); border: 1px solid var(--line); border-radius: 16px; padding: 18px; cursor: pointer; box-shadow: var(--shadow); }
      .feature-card:hover { transform: translateY(-2px); }
      .badge { display: inline-flex; align-items: center; padding: 6px 10px; border-radius: 999px; font-size: 0.75rem; font-weight: 700; }
      .badge.success { background: var(--success-soft); color: var(--success); }
      .badge.warning { background: var(--warning-soft); color: var(--warning); }
      .badge.info { background: var(--primary-soft); color: var(--primary); }
      .doc-header { display: flex; justify-content: space-between; align-items: center; gap: 16px; flex-wrap: wrap; }
      .meta-row { display: flex; gap: 12px; flex-wrap: wrap; }
      .meta { background: var(--soft); border: 1px solid var(--line); border-radius: 10px; padding: 10px 12px; color: var(--muted); font-size: 0.86rem; }
      .actions { display: flex; gap: 10px; flex-wrap: wrap; }
      .section-label { margin: 0 0 14px; letter-spacing: 0.08em; text-transform: uppercase; font-weight: 800; font-size: 0.75rem; color: var(--muted); }
      .summary-card { background: linear-gradient(180deg, #ffffff 0%, #f5f9ff 100%); border: 1px solid var(--line); border-radius: 18px; padding: 22px; }
      .info-grid { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 16px; }
      .mini-card { background: var(--panel); border: 1px solid var(--line); border-radius: 16px; padding: 18px; }
      ul.simple { margin: 10px 0 0; padding-left: 18px; color: var(--muted); }
      .timeline { display: flex; flex-direction: column; gap: 18px; }
      .timeline-item { display: grid; grid-template-columns: 140px 1fr; gap: 16px; align-items: start; }
      .timeline-date { background: var(--soft); border: 1px solid var(--line); border-radius: 10px; padding: 10px 12px; font-weight: 700; color: var(--primary); }
      .timeline-node { position: relative; padding-left: 18px; }
      .timeline-node:before { content: ''; position: absolute; left: 0; top: 8px; width: 10px; height: 10px; background: var(--accent); border-radius: 50%; }
      .timeline-node:after { content: ''; position: absolute; left: 4px; top: 18px; bottom: -18px; width: 2px; background: var(--line); }
      .graph-canvas { background: linear-gradient(180deg, #f9fbff, #f4f7fb); border-radius: 18px; border: 1px solid var(--line); min-height: 260px; padding: 22px; }
      .graph-details { margin-top: 18px; padding-top: 14px; border-top: 1px solid var(--line); color: var(--muted); }
      .graph-root { display: flex; justify-content: center; margin-bottom: 18px; }
      .graph-node { background: white; border: 1px solid var(--line); border-radius: 5px; padding: 10px 14px; box-shadow: none; font-weight: 700; }
      .graph-branch { display: flex; justify-content: center; gap: 18px; flex-wrap: wrap; }
      .graph-node.small { font-size: 0.92rem; }
      .chat-panel { background: var(--panel); border: 1px solid var(--line); border-radius: 18px; box-shadow: var(--shadow); padding: 18px; }
      .chat-list { display: flex; flex-direction: column; gap: 12px; max-height: 300px; overflow: auto; }
      .msg { max-width: 90%; padding: 12px 14px; border-radius: 12px; }
      .msg { white-space: pre-wrap; }
      .evidence-block { padding: 8px 0; }
      .evidence-block + .evidence-block { border-top: 1px solid var(--line); }
      .evidence-label { display: block; margin-bottom: 4px; color: var(--muted); font-size: 10px; font-weight: 800; letter-spacing: .12em; }
      .msg.user { align-self: flex-end; background: var(--primary); color: white; }
      .msg.ai { align-self: flex-start; background: var(--soft); color: var(--text); border: 1px solid var(--line); }
      .chat-compose { display: flex; gap: 10px; margin-top: 18px; }
      .chat-compose input { flex: 1; padding: 12px 14px; border-radius: 10px; border: 1px solid var(--line); background: #fff; }
      .wizard { display: grid; gap: 18px; }
      .wizard-step { background: var(--panel); border: 1px solid var(--line); border-radius: 18px; padding: 18px; }
      .type-grid { display: grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap: 12px; }
      .type-card { border: 1px solid var(--line); border-radius: 12px; padding: 14px; background: var(--soft); cursor: pointer; text-align: center; font-weight: 700; }
      .type-card.selected { background: var(--primary-soft); border-color: rgba(21,58,103,0.3); color: var(--primary); }
      .form-grid { display: grid; grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 14px; }
      .field { display: flex; flex-direction: column; gap: 8px; }
      .field input, .field textarea, .field select { width: 100%; border: 1px solid var(--line); border-radius: 10px; padding: 10px 12px; background: #fff; }
      .field textarea { min-height: 90px; resize: vertical; }
      .empty-state { background: var(--panel); border: 1px dashed var(--line); border-radius: 16px; padding: 32px 18px; text-align: center; color: var(--muted); }
      .inline-status { margin-top: 16px; font-weight: 600; }
      .success { color: var(--success); }
      .error { color: var(--danger); }
      .warning { color: var(--warning); }
      @media (max-width: 1100px) {
        .feature-grid { grid-template-columns: repeat(2, minmax(0, 1fr)); }
        .cards-2, .cards-3, .info-grid, .type-grid { grid-template-columns: 1fr; }
      }
      @media (max-width: 820px) {
        .app-shell { flex-direction: column; }
        .sidebar { width: 100%; border-right: none; border-bottom: 1px solid rgba(255,255,255,0.08); }
        .workspace { padding: 18px; }
        .topbar { flex-direction: column; align-items: flex-start; }
        .chat-compose { flex-direction: column; }
        .form-grid { grid-template-columns: 1fr; }
      }
    </style>
    <style>
      :root {
        --bg: #f8f9f7;
        --panel: #fff;
        --soft: #eef2ef;
        --line: #dce3de;
        --nav: #fff;
        --nav-soft: #eff3f0;
        --text: #192520;
        --muted: #63716a;
        --primary: #173d33;
        --primary-soft: #eaf1ed;
        --accent: #a36a35;
        --success: #287354;
        --success-soft: #e9f3ed;
        --warning: #96611d;
        --warning-soft: #f8f0e4;
        --danger: #a53e3e;
        --danger-soft: #f8ecea;
        --shadow: 0 14px 36px rgba(25, 37, 32, .06);
      }
      html, body { font-family: "Avenir Next", Avenir, "Segoe UI", sans-serif; background: var(--bg); color: var(--text); }
      body { background-image: radial-gradient(#dfe8e2 0.7px, transparent 0.7px); background-size: 22px 22px; }
      .app-shell { display: block; min-height: 100vh; }
      .sidebar { width: 100%; min-height: 70px; padding: 0 5vw; flex-direction: row; align-items: center; gap: 34px; background: rgba(255,255,255,.94); color: var(--text); border: 0; border-bottom: 1px solid var(--line); }
      .brand { padding: 0; margin: 0; border: 0; font: 700 13px/1 "Avenir Next", Avenir, sans-serif; letter-spacing: .14em; white-space: nowrap; color: var(--primary); }
      .nav { flex-direction: row; align-items: center; gap: 4px; }
      .nav-btn, .sidebar-bottom button { width: auto; padding: 10px 13px; border-radius: 6px; color: var(--muted); font-size: 13px; background: transparent; }
      .nav-btn:hover, .sidebar-bottom button:hover { color: var(--primary); background: var(--nav-soft); }
      .nav-btn.active { color: var(--primary); background: var(--primary-soft); box-shadow: none; }
      .sidebar-bottom { margin: 0 0 0 auto; padding: 0; border: 0; flex-direction: row; gap: 4px; }
      .sidebar-bottom button:first-child::before { content: "●"; color: var(--success); font-size: 9px; margin-right: 2px; }
      .workspace { width: min(1180px, 92vw); margin: 0 auto; padding: 30px 0 64px; }
      .topbar { padding: 0 0 22px; margin: 0 0 24px; border: 0; border-bottom: 1px solid var(--line); border-radius: 0; background: transparent; box-shadow: none; backdrop-filter: none; }
      .landing-page .topbar { display: none; }
      .topbar h1 { font: 600 27px/1.2 Georgia, "Times New Roman", serif; letter-spacing: 0; }
      .eyebrow, .section-label { letter-spacing: .12em; }
      .status-pill { border: 0; background: var(--success-soft); color: var(--success); border-radius: 5px; }
      .page.active { animation: arrive .22s ease-out both; }
      @keyframes arrive { from { opacity: .75; transform: translateY(4px); } to { opacity: 1; transform: translateY(0); } }
      #dashboardPage .home-intro { max-width: 760px; padding: 28px 0 30px; }
      #dashboardPage .home-intro .eyebrow { color: var(--success); }
      #dashboardPage .home-intro h1 { margin: 0; max-width: 690px; font: 500 clamp(38px, 5vw, 58px)/1.08 Georgia, "Times New Roman", serif; letter-spacing: 0; }
      #dashboardPage .home-intro p:last-child { max-width: 650px; margin: 18px 0 0; color: var(--muted); font-size: 16px; line-height: 1.65; }
      .upload-zone { position: relative; min-height: 300px; padding: 38px 20px; border: 1px dashed #9eafa4; border-radius: 5px; background: rgba(255,255,255,.88); box-shadow: none; gap: 10px; transition: background .16s, border-color .16s; }
      .upload-zone h3 { font: 500 25px/1.25 Georgia, "Times New Roman", serif; }
      .upload-zone .badge { display: none; }
      .upload-zone .upload-arrow { display: grid; place-items: center; width: 44px; height: 44px; margin-bottom: 5px; border: 1px solid var(--line); border-radius: 50%; color: var(--primary); font-size: 23px; }
      .upload-zone.drag-over { border-color: var(--success); background: #f0f6f1; }
      .upload-zone.drag-over h3 { color: var(--success); }
      .primary-btn, .secondary-btn, .danger-btn, .small-btn { border-radius: 5px; transition: background .16s, color .16s, border-color .16s; }
      .primary-btn { background: var(--primary); }
      .primary-btn:hover { background: #245849; }
      .secondary-btn { background: #fff; }
      button:focus-visible, input:focus-visible, textarea:focus-visible, select:focus-visible, [tabindex="0"]:focus-visible { outline: 3px solid #91b9a4; outline-offset: 3px; }
      .upload-help { margin: 6px 0 0; color: var(--muted); font-size: 12px; }
      .upload-feedback { width: min(100%, 600px); margin: 8px auto 0; padding: 10px 12px; border-radius: 4px; text-align: left; }
      .upload-feedback.error { color: var(--danger); background: var(--danger-soft); }
      .upload-feedback.success { color: var(--success); background: var(--success-soft); }
      .upload-progress { width: min(100%, 600px); margin: 10px auto 0; color: var(--muted); font-size: 13px; text-align: left; }
      .upload-progress p { margin: 6px 0; }
      .capabilities { display: grid; grid-template-columns: repeat(3, minmax(0,1fr)); gap: 30px; padding: 28px 0; border-bottom: 1px solid var(--line); }
      .capability { padding-right: 24px; }
      .capability h3 { margin: 0 0 7px; font: 600 17px Georgia, serif; }
      .capability p { margin: 0; color: var(--muted); font-size: 14px; line-height: 1.55; }
      .card, .mini-card, .chat-panel, .wizard-step, .summary-card { border-radius: 5px; box-shadow: none; }
      .card { border-color: var(--line); }
      .feature-grid { grid-template-columns: repeat(3, minmax(0,1fr)); gap: 0; margin-top: 0; }
      .feature-card { padding: 20px 24px 20px 0; border: 0; border-radius: 0; background: transparent; box-shadow: none; }
      .feature-card:hover { transform: none; }
      .feature-card .badge { display: none; }
      .feature-card h3 { margin-top: 0; font: 600 17px Georgia, serif; }
      .doc-header h2 { font: 500 29px Georgia, serif; }
      .meta, .timeline-date { border-radius: 4px; }
      .graph-canvas { border-radius: 5px; }
      .empty-state { border-radius: 5px; }
      .page-heading { margin: 0 0 20px; }
      .page-heading h2 { margin: 0 0 8px; font: 500 34px Georgia, serif; }
      .page-heading p { margin: 0; color: var(--muted); }
      .current-overview { margin-top: 32px; padding-top: 24px; border-top: 1px solid var(--line); }
      .current-overview h3 { margin: 0 0 8px; font: 500 22px Georgia, serif; }
      .current-overview p { max-width: 780px; line-height: 1.65; }
      .document-visualizations { margin-top: 42px; }
      .workspace-section { padding: 28px 0; border-top: 1px solid var(--line); }
      .section-heading-row { display: flex; justify-content: space-between; align-items: end; margin-bottom: 18px; }
      .section-heading-row h2 { margin: 0; font: 500 24px Georgia, serif; }
      .section-heading-row .section-label { margin-bottom: 6px; }
      .key-terms-grid { display: grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap: 0; }
      .key-term { min-width: 0; padding: 8px 18px 8px 0; border-right: 1px solid var(--line); }
      .key-term + .key-term { padding-left: 18px; }
      .key-term:last-child { border-right: 0; }
      .key-term-value { display: block; margin-bottom: 6px; color: var(--primary); font: 500 22px Georgia, serif; overflow-wrap: anywhere; }
      .key-term-label { color: var(--muted); font-size: 12px; }
      .workspace-timeline { display: grid; grid-template-columns: repeat(auto-fit, minmax(155px, 1fr)); gap: 10px; }
      .timeline-event { min-height: 112px; padding: 14px; border: 1px solid var(--line); border-top: 3px solid var(--accent); background: rgba(255,255,255,.7); text-align: left; cursor: pointer; }
      .timeline-event:hover, .timeline-event.selected { background: var(--primary-soft); }
      .timeline-event-date { display: block; margin-bottom: 9px; color: var(--primary); font-weight: 700; }
      .timeline-event-label { display: block; font-size: 13px; line-height: 1.45; }
      .timeline-event-page, .source-page { display: inline-block; margin-top: 9px; color: var(--muted); font-size: 11px; }
      .timeline-details { min-height: 24px; margin-top: 12px; color: var(--muted); font-size: 13px; }
      .relationship-visual { display: grid; grid-template-columns: 1fr minmax(120px, .6fr) 1fr; align-items: center; gap: 16px; max-width: 760px; }
      .relationship-node { min-height: 90px; padding: 14px; border: 1px solid var(--line); background: #fff; text-align: center; cursor: pointer; }
      .relationship-node:hover, .relationship-node.selected { border-color: var(--primary); background: var(--primary-soft); }
      .relationship-name { display: block; font: 500 19px Georgia, serif; overflow-wrap: anywhere; }
      .relationship-role { display: block; margin-top: 7px; color: var(--muted); font-size: 11px; letter-spacing: .08em; text-transform: uppercase; }
      .relationship-link { text-align: center; color: var(--muted); font-size: 12px; }
      .relationship-property { grid-column: 1 / -1; justify-self: center; width: min(360px, 100%); }
      .obligations-list { display: grid; grid-template-columns: repeat(2, minmax(0,1fr)); gap: 32px; }
      .obligation-party h3 { margin: 0 0 12px; color: var(--primary); font-size: 12px; letter-spacing: .1em; }
      .obligation-item { padding: 11px 0; border-top: 1px solid var(--line); line-height: 1.5; }
      .draft-disclaimer { max-width: 800px; margin: 18px 0; color: var(--muted); font-size: 12px; line-height: 1.5; }
      .answer-content { font-size: 16px; font-weight: 600; line-height: 1.55; }
      .evidence-content { padding: 10px 12px; border-left: 2px solid var(--accent); background: rgba(255,255,255,.75); font: 14px/1.55 Georgia, serif; font-weight: 400; }
      .view-source-btn { margin-top: 7px; padding: 0; border: 0; background: none; color: var(--primary); font-size: 12px; text-decoration: underline; cursor: pointer; }
      .source-highlight { background: #fff4d8; outline: 2px solid #d4a54b; }
      .intelligence-tabs { display: flex; gap: 8px; overflow-x: auto; margin: 22px 0 0; border-bottom: 1px solid var(--line); }
      .intel-tab { padding: 12px 14px; border: 0; border-bottom: 2px solid transparent; background: transparent; color: var(--muted); cursor: pointer; white-space: nowrap; }
      .intel-tab.active { color: var(--primary); border-bottom-color: var(--accent); }
      .intel-tab:focus-visible { outline-offset: -2px; }
      .requires-document[hidden] { display: none; }
      .create-steps { display: flex; gap: 20px; flex-wrap: wrap; margin: 18px 0 24px; color: var(--muted); font-size: 13px; }
      .create-steps strong { color: var(--primary); font-weight: 700; }
      .generated-preview { margin-top: 24px; padding: 26px; border: 1px solid var(--line); background: #fff; }
      .generated-preview pre { max-width: 760px; margin: 0 auto; white-space: pre-wrap; overflow-wrap: anywhere; font: 15px/1.7 Georgia, "Times New Roman", serif; color: #252d29; }
      @media (max-width: 820px) {
        .sidebar { min-height: 0; padding: 16px 4vw; flex-wrap: wrap; gap: 12px; }
        .brand { width: 100%; }
        .nav { overflow-x: auto; max-width: 100%; }
        .nav-btn, .sidebar-bottom button { white-space: nowrap; padding: 9px 10px; }
        .sidebar-bottom { margin-left: 0; }
        .workspace { width: 92vw; padding-top: 22px; }
        .topbar { flex-direction: row; align-items: center; }
        .capabilities, .feature-grid { grid-template-columns: 1fr; gap: 0; }
        .capability { padding: 15px 0; border-bottom: 1px solid var(--line); }
        #dashboardPage .home-intro { padding-top: 14px; }
        #dashboardPage .home-intro h1 { font-size: 40px; }
        .upload-zone { min-height: 270px; }
        .key-terms-grid { grid-template-columns: repeat(2, minmax(0, 1fr)); }
        .key-term { padding: 12px 10px 12px 0; border-right: 0; border-bottom: 1px solid var(--line); }
        .key-term + .key-term { padding-left: 0; }
        .relationship-visual { grid-template-columns: 1fr; }
        .relationship-property { grid-column: auto; width: 100%; }
        .obligations-list { grid-template-columns: 1fr; gap: 20px; }
      }
      @media (max-width: 520px) {
        .sidebar { gap: 8px; }
        .nav { order: 2; width: 100%; }
        .sidebar-bottom { position: absolute; top: 12px; right: 4vw; }
        .sidebar-bottom button { font-size: 0; padding: 8px; }
        .sidebar-bottom button:first-child { font-size: 0; }
        .sidebar-bottom button:first-child::before { font-size: 14px; margin: 0; }
        .workspace { width: 90vw; }
        .topbar { display: block; overflow-wrap: anywhere; }
        .topbar .status-pill { display: inline-block; margin-top: 10px; }
        #dashboardPage .home-intro h1 { font-size: 36px; }
        .upload-zone { padding: 28px 14px; }
        .chat-compose { flex-direction: row; }
        .chat-compose input { min-width: 0; }
      }
    </style>
  </head>
  <body>
    <div class="app-shell">
      <aside class="sidebar">
        <div class="brand">LEGAL AI COPILOT</div>
        <nav class="nav" aria-label="Main navigation">
          <button class="nav-btn active" data-page="dashboard">Home</button>
          <button class="nav-btn" data-page="currentDocument">Documents</button>
          <button class="nav-btn" data-page="askAi">Ask AI</button>
          <button class="nav-btn" data-page="generator">Create Document</button>
        </nav>
        <div class="sidebar-bottom">
          <button data-page="security">Privacy &amp; Security</button>
          <button data-page="settings">Settings</button>
        </div>
      </aside>
      <main class="workspace">
        <header class="topbar">
          <div>
            <p class="eyebrow">AI-Powered Legal Document Intelligence Platform</p>
            <h1 id="pageTitle">Legal Document Intelligence</h1>
          </div>
          <div class="status-pill" id="statusPill">Private Session</div>
        </header>

        <section class="page active" id="dashboardPage">
          <div class="home-intro">
            <p class="eyebrow">A private legal workspace</p>
            <h1>Understand your legal documents with AI.</h1>
            <p>Upload a legal document to extract important information, ask grounded questions, explore relationships, and create new documents.</p>
          </div>
          <div class="upload-zone" id="uploadZone" tabindex="0" role="button" aria-label="Choose or drop a PDF document">
            <div class="upload-arrow" aria-hidden="true">↑</div>
            <h3 id="uploadPrompt">Drop your legal document here</h3>
            <p class="muted" id="uploadSubprompt">Upload a legal document to begin.</p>
            <button class="primary-btn" id="browseUploadBtn" type="button">Browse PDF</button>
            <input id="documentUpload" type="file" accept=".pdf,application/pdf" hidden />
            <p class="upload-help">PDF · Secure temporary session · Up to 15 MB</p>
            <div class="upload-feedback" id="uploadFeedback" role="status" aria-live="polite" hidden></div>
            <div class="upload-progress" id="uploadProgress" aria-live="polite" hidden></div>
          </div>
          <div class="capabilities" aria-label="Document tools">
            <div class="capability"><h3>Understand</h3><p>Extract key information and obligations automatically.</p></div>
            <div class="capability"><h3>Ask</h3><p>Ask questions grounded in the document and its evidence.</p></div>
            <div class="capability"><h3>Create</h3><p>Generate a new legal document from your information.</p></div>
          </div>
        </section>

        <section class="page" id="currentDocumentPage">
          <div class="card">
            <div class="doc-header">
              <div>
                <p class="section-label">Current Document</p>
                <h2 id="currentDocumentName">No document uploaded</h2>
              </div>
              <div class="actions">
                <button class="primary-btn requires-document" id="askAiAction">Ask AI</button>
                <button class="secondary-btn requires-document" id="intelligenceAction">View Intelligence</button>
                <button class="secondary-btn requires-document" id="timelineAction">Timeline</button>
                <button class="secondary-btn requires-document" id="graphAction">Relationships</button>
                <button class="danger-btn requires-document" id="deleteDocumentBtn">Delete Document</button>
              </div>
            </div>
            <div class="meta-row" style="margin-top: 18px;">
              <div class="meta" id="currentDocType">Document type: —</div>
              <div class="meta" id="currentDocStatus">Processing status: Awaiting upload</div>
              <div class="meta" id="currentPageCount">Page count: 0</div>
            </div>
            <div class="inline-status muted" id="currentDocumentStatusText">No document uploaded.</div>
            <div class="current-overview">
              <p class="section-label">Document Overview</p>
              <h3>What is this document about?</h3>
              <p id="currentDocSummary" class="muted">No document yet. Upload a legal document to begin.</p>
              <button class="primary-btn" id="uploadFromEmptyBtn" type="button">Upload Document</button>
            </div>
          </div>
          <div class="document-visualizations">
            <section class="workspace-section" aria-labelledby="keyTermsHeading">
              <div class="section-heading-row"><div><p class="section-label">Extracted terms</p><h2 id="keyTermsHeading">Key terms</h2></div></div>
              <div id="keyTermsContainer" class="key-terms-grid"></div>
            </section>
            <section class="workspace-section" aria-labelledby="workspaceTimelineHeading">
              <div class="section-heading-row"><div><p class="section-label">Dates and obligations</p><h2 id="workspaceTimelineHeading">Agreement timeline</h2></div></div>
              <div id="workspaceTimeline" class="workspace-timeline"></div>
              <div id="workspaceTimelineDetails" class="timeline-details" aria-live="polite"></div>
            </section>
            <section class="workspace-section" aria-labelledby="relationshipHeading">
              <div class="section-heading-row"><div><p class="section-label">Parties and property</p><h2 id="relationshipHeading">Relationships</h2></div></div>
              <div id="workspaceRelationships" class="relationship-visual"></div>
              <div id="workspaceRelationshipDetails" class="graph-details" aria-live="polite"></div>
            </section>
            <section class="workspace-section" aria-labelledby="obligationsHeading">
              <div class="section-heading-row"><div><p class="section-label">By party</p><h2 id="obligationsHeading">Obligations</h2></div></div>
              <div id="workspaceObligations" class="obligations-list"></div>
            </section>
          </div>
        </section>

        <section class="page" id="intelligencePage">
          <div class="page-heading">
            <h2>Document Intelligence</h2>
            <p>Key facts and obligations extracted from the active document.</p>
          </div>
          <div class="card summary-card">
            <p class="section-label">Document Overview</p>
            <h3>What is this document about?</h3>
            <p id="overviewSummary">Document intelligence will appear here after processing.</p>
          </div>
          <nav class="intelligence-tabs" aria-label="Document intelligence sections">
            <button type="button" class="intel-tab active" data-intel-tab="overview">Overview</button>
            <button type="button" class="intel-tab" data-intel-tab="clauses">Clauses</button>
            <button type="button" class="intel-tab" data-intel-tab="obligations">Obligations</button>
            <button type="button" class="intel-tab" data-intel-tab="dates">Dates</button>
            <button type="button" class="intel-tab" data-intel-tab="financials">Financials</button>
          </nav>
          <div class="grid cards-2" style="margin-top: 20px;">
            <div class="mini-card" data-intel-section="overview">
              <p class="section-label">Parties</p>
              <div id="partiesContainer"></div>
            </div>
            <div class="mini-card" data-intel-section="dates">
              <p class="section-label">Important Dates</p>
              <div id="datesContainer"></div>
            </div>
            <div class="mini-card" data-intel-section="dates">
              <p class="section-label">Time Period</p>
              <div id="durationContainer"></div>
            </div>
            <div class="mini-card" data-intel-section="financials">
              <p class="section-label">Financial Terms</p>
              <div id="financialContainer"></div>
            </div>
            <div class="mini-card" data-intel-section="clauses">
              <p class="section-label">Important Clauses</p>
              <div id="clausesContainer"></div>
            </div>
            <div class="mini-card" data-intel-section="obligations">
              <p class="section-label">Obligations</p>
              <div id="obligationsContainer"></div>
            </div>
          </div>
        </section>

        <section class="page" id="timelinePage">
          <div class="page-heading"><h2>Agreement Timeline</h2><p>Events and periods identified in the current document.</p></div>
          <div class="workspace-timeline" id="timelineContainer"></div>
          <div id="timelinePageDetails" class="timeline-details" aria-live="polite"></div>
        </section>

        <section class="page" id="graphPage">
          <div class="page-heading"><h2>Document Relationships</h2><p>Parties, property, and terms extracted from the active agreement.</p></div>
          <div id="graphContainer" class="relationship-visual"></div>
          <div id="graphDetails" class="graph-details" aria-live="polite"></div>
        </section>

        <section class="page" id="askAiPage">
          <div class="chat-panel">
            <p class="section-label">Ask about your document</p>
            <p id="askCurrentDocument" class="current-document-chip">Current document: none</p>
            <div class="grid cards-2" style="margin-bottom: 18px;">
              <button class="secondary-btn small-btn suggested-question" data-question="What is the agreement period?">What is the agreement period?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What is the monthly rent?">What is the monthly rent?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What is the security deposit?">What is the security deposit?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What are the important dates?">What are the important dates?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What are the tenant's obligations?">What are the tenant's obligations?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What is the notice period?">What is the notice period?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="Who is the landlord?">Who is the landlord?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="Who is the tenant?">Who is the tenant?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What is the property address?">What is the property address?</button>
              <button class="secondary-btn small-btn suggested-question" data-question="What information about renewal is present?">What information about renewal is present?</button>
            </div>
            <div class="chat-list" id="chatList"></div>
            <div class="chat-compose">
              <input id="chatInput" type="text" placeholder="Ask a question about your document" />
              <button class="primary-btn" id="sendChatBtn">Send</button>
            </div>
          </div>
        </section>

        <section class="page" id="generatorPage">
          <div class="page-heading">
            <h2>Create a Legal Document</h2>
            <p>Generate a structured draft using your information.</p>
          </div>
          <div class="create-steps" aria-label="Document creation steps">
            <strong>01 Choose type</strong><span>02 Enter details</span><span>03 Review</span><span>04 Download</span>
          </div>
          <div class="wizard">
            <div class="wizard-step">
              <p class="section-label">Step 1</p>
              <h3>Choose document type</h3>
              <div class="type-grid" id="documentTypeGrid">
                <button class="type-card selected" type="button" aria-pressed="true" data-type="Rental Agreement">Rental Agreement</button>
                <button class="type-card" type="button" aria-pressed="false" data-type="Employment Agreement">Employment Agreement</button>
                <button class="type-card" type="button" aria-pressed="false" data-type="NDA">NDA</button>
                <button class="type-card" type="button" aria-pressed="false" data-type="Legal Notice">Legal Notice</button>
                <button class="type-card" type="button" aria-pressed="false" data-type="Service Agreement">Service Agreement</button>
              </div>
            </div>
            <div class="wizard-step">
              <p class="section-label">Step 2</p>
              <h3>Enter information</h3>
              <div class="form-grid">
                <div class="field"><label id="landlordNameLabel" for="landlordName">Landlord full name</label><input id="landlordName" value="Rahul Mehta" /></div>
                <div class="field"><label id="tenantNameLabel" for="tenantName">Tenant full name</label><input id="tenantName" value="Priya Sharma" /></div>
                <div class="field"><label id="landlordAddressLabel" for="landlordAddress">Landlord address</label><input id="landlordAddress" /></div>
                <div class="field"><label id="landlordContactLabel" for="landlordContact">Landlord contact</label><input id="landlordContact" /></div>
                <div class="field"><label id="tenantAddressLabel" for="tenantAddress">Tenant address</label><input id="tenantAddress" /></div>
                <div class="field"><label id="tenantContactLabel" for="tenantContact">Tenant contact</label><input id="tenantContact" /></div>
                <div class="field"><label id="propertyAddressLabel" for="propertyAddress">Property address</label><input id="propertyAddress" value="Flat 402, Green Heights, Andheri East, Mumbai" /></div>
                <div class="field"><label id="propertyTypeLabel" for="propertyType">Property type</label><input id="propertyType" value="Residential Apartment" /></div>
                <div class="field"><label>Start date</label><input id="startDate" type="date" value="2026-09-01" /></div>
                <div class="field"><label>End date</label><input id="endDate" type="date" value="2027-07-31" /></div>
                <div class="field"><label id="monthlyRentLabel" for="monthlyRent">Monthly rent</label><input id="monthlyRent" value="₹30,000" /></div>
                <div class="field"><label id="securityDepositLabel" for="securityDeposit">Security deposit</label><input id="securityDeposit" value="₹60,000" /></div>
                <div class="field"><label id="rentDueDateLabel" for="rentDueDate">Rent due date</label><input id="rentDueDate" value="5th of every month" /></div>
                <div class="field"><label id="noticePeriodLabel" for="noticePeriod">Notice period</label><input id="noticePeriod" value="30 days" /></div>
                <div class="field" style="grid-column: 1 / -1;"><label id="additionalTermsLabel" for="additionalTerms">Additional terms</label><textarea id="additionalTerms">The tenant will maintain the property in good condition and pay all applicable utility charges.</textarea></div>
              </div>
              <div style="margin-top: 18px; display:flex; gap:10px; flex-wrap: wrap;">
                <button class="primary-btn" id="generateDocumentBtn">Generate</button>
                <button class="secondary-btn" id="editGeneratedDocumentBtn" type="button" hidden>Edit</button>
                <a class="secondary-btn" id="downloadPdfBtn" href="#" download hidden>Generate PDF</a>
                <a class="secondary-btn" id="downloadWordBtn" href="#" download hidden>Download Word</a>
              </div>
              <p class="draft-disclaimer">This document is an AI-generated draft for informational and educational purposes only. It is not a substitute for professional legal advice. Please have the document reviewed by a qualified legal professional before signing or using it.</p>
              <div class="generated-preview" id="generatedPreviewPanel" hidden>
                <p class="section-label">Generated document preview</p>
                <article id="generatedPreview" class="generated-document-page"></article>
              </div>
              <div class="inline-status muted" id="generatorStatus">Generate a draft document to create a usable legal template.</div>
            </div>
          </div>
        </section>

        <section class="page" id="securityPage">
          <div class="card">
            <p class="section-label">Privacy &amp; Security</p>
            <div class="badge success">Private Session</div>
            <p style="margin-top: 16px;">Your current document is processed only for this session. Uploading a new document removes the previous document and its associated analysis data.</p>
            <button class="danger-btn" id="deleteFromSecurityBtn">Delete Current Document</button>
          </div>
        </section>

        <section class="page" id="settingsPage">
          <div class="card">
            <p class="section-label">Settings</p>
            <p class="muted">Application settings are intentionally minimal in this MVP to prioritize privacy and active-session document handling.</p>
          </div>
        </section>
      </main>
    </div>

    <script>
      const state = {
        activePage: 'dashboard',
        selectedFile: null,
        uploadStatus: 'idle',
        uploadError: null,
        processingStatus: 'idle',
        currentDocument: null,
        generatedDocument: null,
        generatorType: 'Rental Agreement',
        chat: []
      };

      const pageTitle = document.getElementById('pageTitle');
      const statusPill = document.getElementById('statusPill');

      function setStatus(text, kind = 'info') {
        statusPill.textContent = text;
        statusPill.className = 'status-pill';
        if (kind === 'success') statusPill.classList.add('success');
        if (kind === 'warning') statusPill.classList.add('warning');
        if (kind === 'error') statusPill.classList.add('error');
      }

      function setPage(name) {
        state.activePage = name;
        document.body.classList.toggle('landing-page', name === 'dashboard');
        document.querySelectorAll('.nav-btn').forEach((button) => {
          button.classList.toggle('active', button.dataset.page === name);
        });
        document.querySelectorAll('.page').forEach((page) => {
          page.classList.toggle('active', page.id === (name + 'Page'));
        });
        if (name === 'intelligence') {
          const selectedTab = document.querySelector('.intel-tab.active') || document.querySelector('.intel-tab');
          if (selectedTab) {
            document.querySelectorAll('.intel-tab').forEach((tab) => tab.classList.toggle('active', tab === selectedTab));
            document.querySelectorAll('[data-intel-section]').forEach((section) => {
              section.hidden = section.dataset.intelSection !== selectedTab.dataset.intelTab;
            });
          }
        }

        const labelMap = {
          dashboard: 'Legal Document Intelligence',
          currentDocument: 'Current Document',
          askAi: 'Ask AI',
          intelligence: 'Document Intelligence',
          timeline: 'Timeline & Graph',
          generator: 'Create Document',
          security: 'Privacy & Security',
          settings: 'Settings'
        };
        pageTitle.textContent = labelMap[name] || 'Legal Document Intelligence';
      }

      function updateCurrentDocumentView() {
        const doc = state.currentDocument;
        const nameEl = document.getElementById('currentDocumentName');
        const typeEl = document.getElementById('currentDocType');
        const statusEl = document.getElementById('currentDocStatus');
        const pagesEl = document.getElementById('currentPageCount');
        const statusTextEl = document.getElementById('currentDocumentStatusText');
        if (!doc) {
          nameEl.textContent = 'No document uploaded';
          typeEl.textContent = 'Document type: —';
          statusEl.textContent = 'Processing status: Awaiting upload';
          pagesEl.textContent = 'Page count: 0';
          statusTextEl.textContent = 'No document uploaded.';
          document.getElementById('currentDocSummary').textContent = 'No document yet. Upload a legal document to begin.';
          document.getElementById('uploadFromEmptyBtn').hidden = false;
          document.querySelectorAll('.requires-document').forEach((button) => { button.hidden = true; });
          return;
        }

        nameEl.textContent = doc.filename;
        typeEl.textContent = `Document type: ${doc.analysis.document_type}`;
        statusEl.textContent = 'Processing status: Analyzed successfully';
        pagesEl.textContent = `Page count: ${doc.page_count || 1}`;
        statusTextEl.textContent = '✓ Document analyzed successfully';
        statusTextEl.className = 'inline-status success';
        document.getElementById('currentDocSummary').textContent = doc.analysis.summary;
        document.getElementById('uploadFromEmptyBtn').hidden = true;
        document.querySelectorAll('.requires-document').forEach((button) => { button.hidden = false; });
      }

      function renderOverview() {
        const doc = state.currentDocument;
        const summaryEl = document.getElementById('overviewSummary');
        if (!doc || !doc.analysis) {
          summaryEl.textContent = 'Document intelligence will appear here after processing.';
          return;
        }
        const analysis = doc.analysis;
        summaryEl.textContent = analysis.summary;
        const appendText = (parent, tag, className, text) => {
          const element = document.createElement(tag);
          if (className) element.className = className;
          element.textContent = text;
          parent.append(element);
          return element;
        };

        const partiesContainer = document.getElementById('partiesContainer');
        partiesContainer.replaceChildren();
        (analysis.parties || []).forEach((party) => {
          const item = document.createElement('div');
          item.className = 'fact-row';
          appendText(item, 'strong', '', party.name);
          appendText(item, 'span', 'muted', party.role);
          if (party.address) appendText(item, 'span', 'muted', party.address);
          if (party.contact) appendText(item, 'span', 'muted', party.contact);
          if (party.source_page) appendText(item, 'span', 'source-page', `Page ${party.source_page}`);
          partiesContainer.append(item);
        });
        if (!analysis.parties?.length) appendText(partiesContainer, 'p', 'muted', 'No parties were identified in the document.');

        const dateLabels = {
          'document date': 'Document Date',
          'agreement start': 'Agreement Start',
          'agreement end': 'Agreement End',
          'payment due': 'Rent Due',
          'notice period': 'Notice Period',
          renewal: 'Renewal',
          'other date': 'Other Date'
        };
        const dateItems = analysis.dates?.items || [];
        const datesContainer = document.getElementById('datesContainer');
        datesContainer.replaceChildren();
        const uniqueDates = new Set();
        dateItems.forEach((date) => {
          const key = `${date.type}:${date.date}`;
          if (uniqueDates.has(key)) return;
          uniqueDates.add(key);
          const item = document.createElement('div');
          item.className = 'fact-row';
          appendText(item, 'strong', '', dateLabels[date.type] || date.type);
          appendText(item, 'span', 'muted', date.date);
          appendText(item, 'span', 'source-page', `Page ${date.source_page}`);
          datesContainer.append(item);
        });
        if (!uniqueDates.size) appendText(datesContainer, 'p', 'muted', 'No dates or periods were identified.');

        const durationContainer = document.getElementById('durationContainer');
        durationContainer.textContent = analysis.duration !== 'Not specified' ? analysis.duration : 'Not stated';

        const financialContainer = document.getElementById('financialContainer');
        financialContainer.replaceChildren();
        (analysis.financial_terms || []).forEach((term) => {
          const item = document.createElement('div');
          item.className = 'fact-row';
          appendText(item, 'strong', '', term.type === 'monthly_rent' ? 'Monthly Rent' : 'Security Deposit');
          appendText(item, 'span', 'muted', `${term.display_amount}${term.frequency === 'monthly' ? ' / month' : ''}`);
          if (term.due_date) appendText(item, 'span', 'muted', `Due ${term.due_date}`);
          appendText(item, 'span', 'source-page', `Page ${term.source_page}`);
          financialContainer.append(item);
        });
        if (!analysis.financial_terms?.length) appendText(financialContainer, 'p', 'muted', 'No financial terms were identified.');

        const clausesContainer = document.getElementById('clausesContainer');
        clausesContainer.replaceChildren();
        (analysis.clause_items || []).forEach((clause) => {
          const item = document.createElement('div');
          item.className = 'fact-row';
          appendText(item, 'strong', '', clause.title);
          appendText(item, 'span', 'muted', clause.content);
          appendText(item, 'span', 'source-page', `Page ${clause.source_page}`);
          clausesContainer.append(item);
        });

        const obligationsContainer = document.getElementById('obligationsContainer');
        obligationsContainer.replaceChildren();
        ['Landlord', 'Tenant'].forEach((party) => {
          const heading = appendText(obligationsContainer, 'strong', '', party.toUpperCase());
          const entries = (analysis.obligation_items || []).filter((item) => item.party === party);
          entries.forEach((obligation) => {
            const row = document.createElement('div');
            row.className = 'obligation-item';
            appendText(row, 'span', '', obligation.obligation);
            appendText(row, 'span', 'source-page', `Page ${obligation.source_page}`);
            obligationsContainer.append(row);
          });
          if (!entries.length) appendText(obligationsContainer, 'p', 'muted', `No ${party.toLowerCase()} obligations were explicitly identified.`);
          heading.classList.add('obligation-party-heading');
        });

        renderKeyTerms();
        renderWorkspaceTimeline();
        renderWorkspaceRelationships();
        renderWorkspaceObligations();
      }

      function formatDocumentDate(value) {
        if (!value) return '';
        const months = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December'];
        if (value.includes('-') && value.split('-').length === 3) {
          const [year, month, day] = value.split('-');
          return `${day} ${months[Number(month) - 1] || month} ${year}`;
        }
        const parts = value.replace(',', '').split(' ');
        if (parts.length >= 3 && Number.parseInt(parts[0], 10) && Number.parseInt(parts[2], 10)) {
          const day = String(Number.parseInt(parts[0], 10)).padStart(2, '0');
          return `${day} ${parts[1]} ${parts[2]}`;
        }
        return value;
      }

      function renderKeyTerms() {
        const container = document.getElementById('keyTermsContainer');
        const analysis = state.currentDocument?.analysis;
        container.replaceChildren();
        if (!analysis) return;
        const terms = analysis.key_terms || {};
        const items = [
          { label: 'Term', value: terms.term },
          { label: 'Monthly Rent', value: terms.rent ? `${terms.rent.display_amount} / month` : null },
          { label: 'Security Deposit', value: terms.deposit?.display_amount },
          { label: 'Notice', value: terms.notice?.duration ? `${terms.notice.duration} ${terms.notice.unit}` : null },
          { label: 'Rent Due', value: terms.rent_due }
        ].filter((item) => item.value);
        items.forEach((item) => {
          const term = document.createElement('div');
          term.className = 'key-term';
          const value = document.createElement('span');
          value.className = 'key-term-value';
          value.textContent = item.value;
          const label = document.createElement('span');
          label.className = 'key-term-label';
          label.textContent = item.label;
          term.append(value, label);
          container.append(term);
        });
        if (!items.length) container.textContent = 'No key terms were identified in this document.';
      }

      function timelineLabel(type) {
        return ({
          'agreement start': 'Agreement Start',
          'agreement end': 'Agreement End',
          'payment due': 'Rent Due',
          'notice period': 'Notice Period',
          renewal: 'Renewal',
          'document date': 'Document Date'
        })[type] || type;
      }

      function renderWorkspaceTimeline(containerId = 'workspaceTimeline', detailsId = 'workspaceTimelineDetails') {
        const container = document.getElementById(containerId);
        const details = document.getElementById(detailsId);
        const analysis = state.currentDocument?.analysis;
        container.replaceChildren();
        details.textContent = '';
        if (!analysis) return;
        const order = {'document date': 0, 'agreement start': 1, 'payment due': 2, 'notice period': 3, renewal: 4, 'agreement end': 5};
        const events = (analysis.timeline || []).filter((event) => Object.prototype.hasOwnProperty.call(order, event.event)).sort((a, b) => order[a.event] - order[b.event]);
        events.forEach((event) => {
          const button = document.createElement('button');
          button.type = 'button';
          button.className = 'timeline-event';
          const date = document.createElement('span');
          date.className = 'timeline-event-date';
          date.textContent = event.event === 'agreement start' || event.event === 'agreement end' || event.event === 'document date' ? formatDocumentDate(event.date) : event.date;
          const label = document.createElement('span');
          label.className = 'timeline-event-label';
          label.textContent = timelineLabel(event.event);
          const page = document.createElement('span');
          page.className = 'timeline-event-page';
          page.textContent = `Page ${event.page}`;
          button.append(date, label, page);
          button.addEventListener('click', () => {
            container.querySelectorAll('.timeline-event').forEach((item) => item.classList.toggle('selected', item === button));
            details.textContent = `${event.description || event.source_text} · Page ${event.page}`;
          });
          container.append(button);
        });
        if (!events.length) container.textContent = 'No dated events or periods were found in the document.';
      }

      function renderWorkspaceRelationships(containerId = 'workspaceRelationships', detailsId = 'workspaceRelationshipDetails') {
        const container = document.getElementById(containerId);
        const details = document.getElementById(detailsId);
        const analysis = state.currentDocument?.analysis;
        container.replaceChildren();
        details.textContent = '';
        if (!analysis) return;
        const parties = analysis.parties || [];
        const addNode = (name, role, sourcePage, sourceText, extraClass = '') => {
          const node = document.createElement('button');
          node.type = 'button';
          node.className = `relationship-node ${extraClass}`.trim();
          const nodeName = document.createElement('span');
          nodeName.className = 'relationship-name';
          nodeName.textContent = name;
          const nodeRole = document.createElement('span');
          nodeRole.className = 'relationship-role';
          nodeRole.textContent = role;
          node.append(nodeName, nodeRole);
          node.addEventListener('click', () => {
            container.querySelectorAll('.relationship-node').forEach((item) => item.classList.toggle('selected', item === node));
            details.textContent = `${name} · ${role}${sourcePage ? ` · Page ${sourcePage}` : ''}${sourceText ? ` · ${sourceText}` : ''}`;
          });
          container.append(node);
        };
        if (parties.length) {
          addNode(parties[0].name, parties[0].role, parties[0].source_page, parties[0].source_text);
          const link = document.createElement('div');
          link.className = 'relationship-link';
          link.textContent = analysis.document_type;
          container.append(link);
          if (parties[1]) addNode(parties[1].name, parties[1].role, parties[1].source_page, parties[1].source_text);
        }
        const property = analysis.property || {};
        if (property.address) addNode(property.address, 'Property', property.source_page, property.source_text, 'relationship-property');
        if (!parties.length && !property.address) container.textContent = 'No named parties or property were extracted.';
      }

      function renderWorkspaceObligations() {
        const container = document.getElementById('workspaceObligations');
        const analysis = state.currentDocument?.analysis;
        container.replaceChildren();
        if (!analysis) return;
        ['Landlord', 'Tenant'].forEach((party) => {
          const group = document.createElement('div');
          group.className = 'obligation-party';
          const heading = document.createElement('h3');
          heading.textContent = party.toUpperCase();
          group.append(heading);
          const items = (analysis.obligation_items || []).filter((item) => item.party === party);
          items.forEach((item) => {
            const row = document.createElement('div');
            row.className = 'obligation-item';
            row.textContent = item.obligation;
            const page = document.createElement('span');
            page.className = 'source-page';
            page.textContent = `Page ${item.source_page}`;
            row.append(page);
            group.append(row);
          });
          if (!items.length) {
            const empty = document.createElement('p');
            empty.className = 'muted';
            empty.textContent = 'No explicit obligations found.';
            group.append(empty);
          }
          container.append(group);
        });
      }

      function renderTimeline() {
        renderWorkspaceTimeline('timelineContainer', 'timelinePageDetails');
      }

      function renderGraph() {
        renderWorkspaceRelationships('graphContainer', 'graphDetails');
      }

      function showGraphDetails(nodeId) {
        const doc = state.currentDocument;
        if (!doc || !doc.analysis || !doc.analysis.graph) return;
        const node = doc.analysis.graph.nodes.find((item) => item.id === nodeId);
        const details = document.getElementById('graphDetails');
        if (!node) {
          details.textContent = 'No graph details available.';
          return;
        }
        details.replaceChildren();
        const name = document.createElement('strong');
        name.textContent = node.label;
        const role = document.createElement('p');
        role.className = 'muted';
        role.textContent = `Role: ${node.type}`;
        const related = document.createElement('p');
        related.className = 'muted';
        related.textContent = `Related clauses: ${(doc.analysis.clauses || []).slice(0, 3).join(', ')}`;
        details.append(name, role, related);
        if (node.source_page) {
          const source = document.createElement('p');
          source.className = 'muted';
          source.textContent = `Source: Page ${node.source_page}${node.details ? ` · ${node.details}` : ''}`;
          details.append(source);
        }
      }

      function renderChat() {
        const chatList = document.getElementById('chatList');
        chatList.replaceChildren();
        if (!state.currentDocument) {
          const empty = document.createElement('div');
          empty.className = 'empty-state';
          empty.textContent = 'Ask a question about your uploaded document.';
          chatList.append(empty);
          return;
        }
        if (!state.chat.length) {
          const empty = document.createElement('div');
          empty.className = 'empty-state';
          empty.textContent = 'Ask a question about your uploaded document.';
          chatList.append(empty);
          return;
        }
        state.chat.forEach((message) => {
          const item = document.createElement('div');
          item.className = `msg ${message.role === 'user' ? 'user' : 'ai'}`;
          const newline = String.fromCharCode(10);
          const sections = message.text.split(newline + newline);
          const labels = ['ANSWER', 'SOURCE', 'EVIDENCE'];
          const isStructuredAnswer = message.role === 'ai' && sections.every((section) => labels.includes(section.slice(0, section.indexOf(newline))));
          if (isStructuredAnswer) {
            sections.forEach((section) => {
              const separator = section.indexOf(newline);
              const block = document.createElement('div');
              block.className = 'evidence-block';
              const label = document.createElement('span');
              label.className = 'evidence-label';
              label.textContent = section.slice(0, separator);
              const content = document.createElement('div');
              content.textContent = section.slice(separator + 1).trim();
              block.append(label, content);
              item.append(block);
            });
          } else {
            item.textContent = message.text;
          }
          chatList.append(item);
        });
        chatList.scrollTop = chatList.scrollHeight;
      }

      function sendQuestion(question) {
        if (!state.currentDocument) {
          setStatus('Upload a document first', 'warning');
          return;
        }
        const text = (question || document.getElementById('chatInput').value || '').trim();
        if (!text) {
          return;
        }
        document.getElementById('chatInput').value = '';
        state.chat.push({ role: 'user', text });
        renderChat();
        setStatus('Searching document...', 'info');

        fetch('/api/documents/current/chat', {
          method: 'POST',
          headers: { 'Content-Type': 'application/json' },
          body: JSON.stringify({ question: text })
        })
          .then((res) => res.json())
          .then((payload) => {
            const answer = payload.answer || 'Unable to find relevant information in the current document.';
            state.chat.push({ role: 'ai', text: `ANSWER\n${answer}\n\nSOURCE\n${payload.source || 'Page 1'}\n\nEVIDENCE\n${payload.evidence || 'No evidence found.'}` });
            renderChat();
            setStatus('Answer ready', 'success');
          })
          .catch(() => {
            state.chat.push({ role: 'ai', text: 'Something went wrong. Please try again.' });
            renderChat();
            setStatus('Unable to answer', 'error');
          });
      }

      function hydrateCurrentDocument() {
        fetch('/api/documents/current')
          .then((response) => response.json())
          .then((payload) => {
            if (payload && payload.filename) {
              state.currentDocument = payload;
              updateCurrentDocumentView();
              renderOverview();
              renderTimeline();
              renderGraph();
              renderChat();
            }
          })
          .catch(() => {});
      }

      function clearUploadFeedback() {
        state.uploadError = null;
        const feedback = document.getElementById('uploadFeedback');
        feedback.hidden = true;
        feedback.textContent = '';
        feedback.className = 'upload-feedback';
      }

      function showUploadError(message) {
        state.uploadError = message;
        state.uploadStatus = 'idle';
        state.processingStatus = 'idle';
        const feedback = document.getElementById('uploadFeedback');
        feedback.textContent = message;
        feedback.className = 'upload-feedback error';
        feedback.hidden = false;
        document.getElementById('uploadProgress').hidden = true;
      }

      function resetUploadExperience() {
        state.selectedFile = null;
        state.uploadStatus = 'idle';
        state.uploadError = null;
        state.processingStatus = 'idle';
        document.getElementById('documentUpload').value = '';
        document.getElementById('uploadPrompt').textContent = 'Drop your legal document here';
        document.getElementById('uploadSubprompt').textContent = 'Upload a legal document to begin.';
        document.getElementById('uploadFeedback').hidden = true;
        document.getElementById('uploadProgress').hidden = true;
      }

      function renderUploadProgress(complete) {
        const progress = document.getElementById('uploadProgress');
        const steps = ['Uploading document', 'Extracting text', 'Analyzing document', 'Preparing document intelligence'];
        progress.innerHTML = steps.map((step, index) => `<p>${complete ? '✓' : index === 0 ? '●' : '○'} ${step}${!complete && index === 0 ? '...' : ''}</p>`).join('');
        progress.hidden = false;
      }

      function uploadErrorMessage(code) {
        const messages = {
          invalid: 'Please upload a valid PDF document.',
          invalid_file: 'Please upload a valid PDF document.',
          too_large: 'File is too large. Please upload a smaller PDF.',
          processing_failed: "We couldn't process this document. Please try another PDF.",
          ocr_failed: "We couldn't read this scanned document.",
          no_readable_text: 'No readable text was found in this document.',
          analysis_failed: 'Unable to analyze this document right now. Please try again.'
        };
        return messages[code] || 'Unable to upload the document. Please try again.';
      }

      async function handleUpload(file) {
        if (!file) return;
        state.selectedFile = file;
        clearUploadFeedback();
        const isPdfName = file.name.toLowerCase().endsWith('.pdf');
        const hasInvalidMime = file.type && file.type.toLowerCase() !== 'application/pdf';
        if (!isPdfName || hasInvalidMime) {
          showUploadError(uploadErrorMessage('invalid'));
          return;
        }
        if (file.size > 15 * 1024 * 1024) {
          showUploadError(uploadErrorMessage('too_large'));
          return;
        }

        state.uploadStatus = 'uploading';
        state.processingStatus = 'pending';
        document.getElementById('uploadPrompt').textContent = file.name;
        document.getElementById('uploadSubprompt').textContent = 'Uploading document...';
        document.getElementById('browseUploadBtn').disabled = true;
        renderUploadProgress(false);
        const formData = new FormData();
        formData.append('file', file);

        try {
          const response = await fetch('/api/documents/upload', { method: 'POST', body: formData });
          const payload = await response.json();
          if (!response.ok || !payload.ok) {
            showUploadError(uploadErrorMessage(payload.error));
            document.getElementById('uploadPrompt').textContent = 'Drop your legal document here';
            document.getElementById('uploadSubprompt').textContent = 'Upload a legal document to begin.';
            return;
          }
          state.uploadStatus = 'complete';
          state.processingStatus = 'complete';
          state.currentDocument = payload.document;
          state.chat = [];
          renderUploadProgress(true);
          document.getElementById('uploadSubprompt').textContent = 'Your document is ready.';
          const feedback = document.getElementById('uploadFeedback');
          feedback.textContent = 'Document uploaded.';
          feedback.className = 'upload-feedback success';
          feedback.hidden = false;
          setPage('currentDocument');
          updateCurrentDocumentView();
          renderOverview();
          renderTimeline();
          renderGraph();
          renderChat();
          setStatus('Private Session', 'success');
        } catch {
          showUploadError(uploadErrorMessage('network'));
          document.getElementById('uploadPrompt').textContent = 'Drop your legal document here';
          document.getElementById('uploadSubprompt').textContent = 'Upload a legal document to begin.';
        } finally {
          state.uploadStatus = state.uploadStatus === 'complete' ? 'complete' : 'idle';
          document.getElementById('browseUploadBtn').disabled = false;
        }
      }

      function deleteCurrentDocument() {
        if (!state.currentDocument) return;
        const confirmed = window.confirm('Delete this document and all associated analysis?');
        if (!confirmed) return;
        fetch('/api/documents/current', { method: 'DELETE' })
          .then((response) => response.json())
          .then((payload) => {
            state.currentDocument = null;
            state.chat = [];
            resetUploadExperience();
            renderOverview();
            renderTimeline();
            renderGraph();
            renderChat();
            updateCurrentDocumentView();
            setPage('dashboard');
            setStatus('Document securely removed', 'success');
          })
          .catch(() => {
            setStatus('Unable to delete the current document.', 'error');
          });
      }

      function renderGeneratorStatus(message, kind = 'muted') {
        const el = document.getElementById('generatorStatus');
        el.textContent = message;
        el.className = `inline-status ${kind}`;
      }

      function renderGeneratedPreview(blocks) {
        const preview = document.getElementById('generatedPreview');
        preview.replaceChildren();
        blocks.forEach((block) => {
          const tag = block.type === 'title' ? 'h2' : block.type === 'section_heading' ? 'h3' : block.type === 'party_heading' ? 'h4' : 'p';
          const element = document.createElement(tag);
          element.className = `preview-${block.type}`;
          element.textContent = block.text;
          preview.append(element);
        });
      }

      document.querySelectorAll('.nav-btn').forEach((button) => {
        button.addEventListener('click', () => {
          setPage(button.dataset.page);
        });
      });

      document.querySelectorAll('.sidebar-bottom [data-page]').forEach((button) => {
        button.addEventListener('click', () => setPage(button.dataset.page));
      });

      document.querySelectorAll('.feature-card').forEach((card) => {
        card.addEventListener('click', () => {
          setPage(card.dataset.page);
        });
      });

      document.getElementById('browseUploadBtn').addEventListener('click', (event) => {
        event.stopPropagation();
        const input = document.getElementById('documentUpload');
        state.selectedFile = null;
        input.value = '';
        input.click();
      });

      document.getElementById('documentUpload').addEventListener('change', (event) => {
        handleUpload(event.target.files[0]);
      });

      const uploadZone = document.getElementById('uploadZone');
      uploadZone.addEventListener('click', (event) => {
        if (event.target.closest('button')) return;
        document.getElementById('browseUploadBtn').click();
      });
      uploadZone.addEventListener('keydown', (event) => {
        if (event.key === 'Enter' || event.key === ' ') {
          event.preventDefault();
          document.getElementById('browseUploadBtn').click();
        }
      });
      uploadZone.addEventListener('dragover', (event) => {
        event.preventDefault();
        uploadZone.classList.add('drag-over');
        document.getElementById('uploadPrompt').textContent = 'Drop your PDF here';
      });
      uploadZone.addEventListener('dragleave', (event) => {
        if (!uploadZone.contains(event.relatedTarget)) {
          uploadZone.classList.remove('drag-over');
          document.getElementById('uploadPrompt').textContent = 'Drop your legal document here';
        }
      });
      uploadZone.addEventListener('drop', (event) => {
        event.preventDefault();
        uploadZone.classList.remove('drag-over');
        handleUpload(event.dataTransfer.files[0]);
      });

      document.getElementById('askAiAction').addEventListener('click', () => setPage('askAi'));
      document.getElementById('intelligenceAction').addEventListener('click', () => setPage('intelligence'));
      document.getElementById('timelineAction').addEventListener('click', () => setPage('timeline'));
      document.getElementById('graphAction').addEventListener('click', () => setPage('graph'));
      document.getElementById('uploadFromEmptyBtn').addEventListener('click', () => {
        setPage('dashboard');
        document.getElementById('browseUploadBtn').click();
      });
      document.getElementById('deleteDocumentBtn').addEventListener('click', deleteCurrentDocument);
      document.getElementById('deleteFromSecurityBtn').addEventListener('click', deleteCurrentDocument);

      document.querySelectorAll('.intel-tab').forEach((button) => {
        button.addEventListener('click', () => {
          document.querySelectorAll('.intel-tab').forEach((tab) => tab.classList.toggle('active', tab === button));
          document.querySelectorAll('[data-intel-section]').forEach((section) => {
            section.hidden = section.dataset.intelSection !== button.dataset.intelTab;
          });
        });
      });

      document.querySelectorAll('.suggested-question').forEach((button) => {
        button.addEventListener('click', () => sendQuestion(button.dataset.question));
      });

      document.getElementById('sendChatBtn').addEventListener('click', () => sendQuestion(null));
      document.getElementById('chatInput').addEventListener('keydown', (event) => {
        if (event.key === 'Enter') sendQuestion(null);
      });

      document.querySelectorAll('.type-card').forEach((card) => {
        card.addEventListener('click', () => {
          document.querySelectorAll('.type-card').forEach((item) => item.classList.remove('selected'));
          card.classList.add('selected');
          document.querySelectorAll('.type-card').forEach((item) => item.setAttribute('aria-pressed', String(item === card)));
          state.generatorType = card.dataset.type;
          const formLabels = {
            'Rental Agreement': ['Landlord full name', 'Tenant full name', 'Landlord address', 'Landlord contact', 'Tenant address', 'Tenant contact', 'Property address', 'Property type', 'Monthly rent', 'Security deposit', 'Rent due date', 'Notice period'],
            'Employment Agreement': ['Employer name', 'Employee name', 'Employer address', 'Employer contact', 'Employee address', 'Employee contact', 'Work location', 'Position or role', 'Compensation', 'Other compensation', 'Pay schedule', 'Notice period'],
            'NDA': ['Disclosing party', 'Receiving party', 'Party address', 'Party contact', 'Receiving party address', 'Receiving party contact', 'Purpose address', 'Purpose', 'Fee, if applicable', 'Security, if applicable', 'Effective date', 'Notice period'],
            'Legal Notice': ['From', 'To', 'Sender address', 'Sender contact', 'Recipient address', 'Recipient contact', 'Notice address', 'Subject', 'Amount in dispute', 'Advance paid', 'Response date', 'Response period'],
            'Service Agreement': ['Service provider', 'Client', 'Provider address', 'Provider contact', 'Client address', 'Client contact', 'Service location', 'Services', 'Service fee', 'Advance payment', 'Payment due', 'Notice period']
          }[state.generatorType];
          const labelIds = ['landlordNameLabel', 'tenantNameLabel', 'landlordAddressLabel', 'landlordContactLabel', 'tenantAddressLabel', 'tenantContactLabel', 'propertyAddressLabel', 'propertyTypeLabel', 'monthlyRentLabel', 'securityDepositLabel', 'rentDueDateLabel', 'noticePeriodLabel'];
          labelIds.forEach((id, index) => { document.getElementById(id).textContent = formLabels[index]; });
        });
      });

      document.getElementById('generateDocumentBtn').addEventListener('click', () => {
        const payload = {
          documentType: state.generatorType,
          landlordName: document.getElementById('landlordName').value,
          tenantName: document.getElementById('tenantName').value,
          landlordAddress: document.getElementById('landlordAddress').value,
          landlordContact: document.getElementById('landlordContact').value,
          tenantAddress: document.getElementById('tenantAddress').value,
          tenantContact: document.getElementById('tenantContact').value,
          propertyAddress: document.getElementById('propertyAddress').value,
          propertyType: document.getElementById('propertyType').value,
          startDate: document.getElementById('startDate').value,
          endDate: document.getElementById('endDate').value,
          monthlyRent: document.getElementById('monthlyRent').value,
          securityDeposit: document.getElementById('securityDeposit').value,
          rentDueDate: document.getElementById('rentDueDate').value,
          noticePeriod: document.getElementById('noticePeriod').value,
          additionalTerms: document.getElementById('additionalTerms').value
        };

        renderGeneratorStatus('Preparing legal document...', 'warning');
        fetch('/api/documents/generate', {
          method: 'POST',
          headers: { 'Content-Type': 'application/json' },
          body: JSON.stringify(payload)
        })
          .then((response) => response.json())
          .then((result) => {
            if (!result.ok) throw new Error(result.error || 'Generation failed');
            state.generatedDocument = result.document;
            renderGeneratedPreview(result.document.document_blocks || [{ type: 'paragraph', text: result.document.preview_text || result.document.title }]);
            document.getElementById('generatedPreviewPanel').hidden = false;
            document.getElementById('editGeneratedDocumentBtn').hidden = false;
            renderGeneratorStatus('✓ Document generated successfully', 'success');
            document.getElementById('downloadPdfBtn').href = `/api/documents/generated/${result.document.id}/pdf`;
            document.getElementById('downloadWordBtn').href = `/api/documents/generated/${result.document.id}/docx`;
            document.getElementById('downloadPdfBtn').hidden = false;
            document.getElementById('downloadWordBtn').hidden = false;
            setStatus('Document generated successfully', 'success');
          })
          .catch(() => {
            renderGeneratorStatus('Something went wrong. Please try again.', 'error');
            setStatus('Document generation failed', 'error');
          });
      });

      document.getElementById('editGeneratedDocumentBtn').addEventListener('click', () => {
        document.getElementById('landlordName').scrollIntoView({ behavior: 'smooth', block: 'center' });
        document.getElementById('landlordName').focus();
      });

      setPage('dashboard');
      updateCurrentDocumentView();
      renderOverview();
      renderTimeline();
      renderGraph();
      renderChat();
      hydrateCurrentDocument();
    </script>
  </body>
</html>
"""


class LegalAICopilotHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == '/':
            self.send_html(INDEX_HTML)
            return
        if parsed.path.startswith('/api/documents/generated/'):
            return self.handle_generated_document(parsed.path)
        if parsed.path == '/api/documents/current':
            return self.handle_get_current_document()
        if parsed.path == '/api/documents/current/analysis':
            return self.handle_get_analysis()
        if parsed.path == '/api/documents/current/graph':
            return self.handle_get_graph()
        if parsed.path == '/api/documents/current/timeline':
            return self.handle_get_timeline()
        self.send_json(404, {'error': 'Not found'})

    def do_POST(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/documents/upload':
            return self.handle_upload_document()
        if parsed.path == '/api/documents/current/chat':
            return self.handle_chat()
        if parsed.path == '/api/documents/generate':
            return self.handle_generate_document()
        self.send_json(404, {'error': 'Not found'})

    def do_DELETE(self):
        parsed = urlparse(self.path)
        if parsed.path == '/api/documents/current':
            return self.handle_delete_current_document()
        self.send_json(404, {'error': 'Not found'})

    def handle_upload_document(self):
          try:
              ensure_storage_dirs()
              length = int(self.headers.get('Content-Length', '0'))
              content_type = self.headers.get('Content-Type', '')
              if length > MAX_UPLOAD_SIZE + 64 * 1024:
                  return self.send_json(413, {'ok': False, 'error': 'too_large', 'message': 'File is too large.'})
              body = self.rfile.read(length)
              if not body:
                  return self.send_json(400, {'ok': False, 'error': 'invalid_file', 'message': 'No file was uploaded.'})

              if 'multipart/form-data' not in content_type:
                  return self.send_json(400, {'ok': False, 'error': 'invalid_file', 'message': 'Please upload a valid PDF document.'})

              message = BytesParser(policy=policy.default).parsebytes(
                  f'Content-Type: {content_type}\r\n\r\n'.encode('utf-8') + body
              )
              file_part = next((part for part in message.iter_parts() if part.get_filename()), None)
              if file_part is None:
                  return self.send_json(400, {'ok': False, 'error': 'invalid_file', 'message': 'No file was uploaded.'})

              uploaded_file = file_part.get_payload(decode=True) or b''
              filename = os.path.basename(file_part.get_filename())
              part_content_type = file_part.get('Content-Type')
              validation_error = validate_pdf_upload(filename, part_content_type, uploaded_file)
              if validation_error:
                  status = 413 if validation_error == 'too_large' else 400
                  message_text = 'File is too large.' if validation_error == 'too_large' else 'Please upload a valid PDF document.'
                  return self.send_json(status, {'ok': False, 'error': validation_error, 'message': message_text})

              clear_active_document()
              file_id = str(uuid.uuid4())
              target_path = UPLOAD_DIR / f'{file_id}_{filename}'
              target_path.write_bytes(uploaded_file)

              extracted = extract_pdf_text(str(target_path))
              if not extracted.get('ok'):
                  target_path.unlink(missing_ok=True)
                  return self.send_json(422, {'ok': False, 'error': 'processing_failed', 'message': "We couldn't process this document. Please try another PDF."})
              if not extracted.get('text', '').strip():
                  target_path.unlink(missing_ok=True)
                  scanned = any(page.get('has_images') for page in extracted.get('pages', []))
                  code = 'ocr_failed' if scanned else 'no_readable_text'
                  message_text = "We couldn't read this scanned document." if scanned else 'No readable text was found in this document.'
                  return self.send_json(422, {'ok': False, 'error': code, 'message': message_text})
              try:
                  analysis = analyze_document(filename, extracted['text'], pages=extracted['pages'], document_id=file_id)
              except Exception:
                  target_path.unlink(missing_ok=True)
                  return self.send_json(503, {'ok': False, 'error': 'analysis_failed', 'message': 'Unable to analyze this document right now. Please try again.'})
              SESSION['active_document'] = {
                  'id': file_id,
                  'filename': filename,
                  'file_path': str(target_path),
                  'page_count': len(extracted.get('pages', [])) or 1,
                  'analysis': analysis,
                  'created_at': datetime.now(timezone.utc).isoformat(),
              }
              return self.send_json(200, {'ok': True, 'document': SESSION['active_document']})
          except Exception:
              logging.exception('Document upload failed')
              return self.send_json(500, {'ok': False, 'error': 'upload_failed', 'message': 'Unable to upload the document. Please try again.'})

    def handle_get_current_document(self):
        doc = SESSION.get('active_document')
        if not doc:
            return self.send_json(200, {'filename': None, 'analysis': None})
        return self.send_json(200, doc)

    def handle_get_analysis(self):
        doc = SESSION.get('active_document')
        if not doc:
            return self.send_json(200, {'summary': 'Document intelligence will appear here after processing.'})
        return self.send_json(200, doc['analysis'])

    def handle_get_graph(self):
        doc = SESSION.get('active_document')
        if not doc:
            return self.send_json(200, {'nodes': [], 'edges': []})
        return self.send_json(200, doc['analysis'].get('graph', {'nodes': [], 'edges': []}))

    def handle_get_timeline(self):
        doc = SESSION.get('active_document')
        if not doc:
            return self.send_json(200, {'timeline': []})
        return self.send_json(200, {'timeline': doc['analysis'].get('timeline', [])})

    def handle_chat(self):
        payload = self.read_json_body()
        question = str(payload.get('question', '')).strip()
        doc = SESSION.get('active_document')
        if not doc:
            return self.send_json(400, {'error': 'No document is currently active.'})
        answer = answer_question_from_document(question, doc['analysis'], document_id=doc['id'])
        SESSION['chat_history'].append({'role': 'user', 'text': question})
        SESSION['chat_history'].append({'role': 'ai', 'text': answer['answer']})
        return self.send_json(200, answer)

    def handle_generate_document(self):
        payload = self.read_json_body()
        document = build_document_template(payload)
        file_id = document['id']
        doc_text = document['document_text']
        document_blocks = document['document_blocks']
        pdf_bytes = build_pdf_document(document['title'], document_blocks)
        docx_bytes = build_docx_document(document['title'], document_blocks)
        pdf_path = GENERATED_DIR / f'{file_id}.pdf'
        docx_path = GENERATED_DIR / f'{file_id}.docx'
        pdf_path.write_bytes(pdf_bytes)
        docx_path.write_bytes(docx_bytes)
        register_generated_document(file_id, {'id': file_id, 'title': document['title'], 'pdf_path': str(pdf_path), 'docx_path': str(docx_path), 'disclaimer': document['disclaimer']})
        return self.send_json(200, {'ok': True, 'document': {'id': file_id, 'title': document['title'], 'preview_text': doc_text, 'disclaimer': document['disclaimer'], 'download_pdf': f'/api/documents/generated/{file_id}/pdf', 'download_docx': f'/api/documents/generated/{file_id}/docx'}})

    def handle_generated_document(self, path):
        matches = re.search(r'/api/documents/generated/([^/]+)/(pdf|docx)$', path)
        if not matches:
            return self.send_json(404, {'error': 'Not found'})
        doc_id, file_type = matches.groups()
        generated = SESSION['generated_documents'].get(doc_id)
        if not generated:
            return self.send_json(404, {'error': 'Document not found'})
        target_path = generated['pdf_path'] if file_type == 'pdf' else generated['docx_path']
        if not os.path.exists(target_path):
            return self.send_json(404, {'error': 'Document not found'})
        with open(target_path, 'rb') as handle:
            content = handle.read()
        content_type = 'application/pdf' if file_type == 'pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        self.send_response(200)
        self.send_header('Content-Type', content_type)
        self.send_header('Content-Disposition', f'attachment; filename="{doc_id}.{file_type}"')
        self.send_header('Content-Length', str(len(content)))
        self.end_headers()
        self.wfile.write(content)
        return None

    def handle_delete_current_document(self):
        clear_active_document()
        return self.send_json(200, {'ok': True, 'message': 'Document securely removed.'})

    def read_json_body(self):
        try:
            length = int(self.headers.get('Content-Length', '0'))
            raw = self.rfile.read(length)
            if not raw:
                return {}
            return json.loads(raw.decode('utf-8'))
        except Exception:
            return {}

    def send_html(self, html):
        body = html.encode('utf-8')
        self.send_response(200)
        self.send_header('Content-Type', 'text/html; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, format, *args):
        return


def main():
    ensure_storage_dirs()
    host = '127.0.0.1'
    port = 8000
    print(f'Legal AI Copilot running at http://{host}:{port}')
    HTTPServer((host, port), LegalAICopilotHandler).serve_forever()


if __name__ == '__main__':
    main()
